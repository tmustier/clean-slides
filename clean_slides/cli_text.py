"""Text argument parsing and writing helpers for CLI edit/batch commands."""

from __future__ import annotations

import json
from typing import TYPE_CHECKING

from pptx.shapes.autoshape import Shape
from typing_extensions import TypeGuard

if TYPE_CHECKING:
    from .editor import ParagraphSpec

RunOverridesMap = dict[str, object]
TextRun = tuple[str, RunOverridesMap]


def _is_object_list(value: object) -> TypeGuard[list[object]]:
    return isinstance(value, list)


def _is_str_object_dict(value: object) -> TypeGuard[dict[str, object]]:
    return isinstance(value, dict)


def _coerce_overrides(value: object) -> RunOverridesMap:
    if _is_str_object_dict(value):
        return dict(value)
    return {}


def _expand_newlines(result: list[TextRun], text: str, opts: RunOverridesMap) -> None:
    """Split text on newlines, inserting line break markers."""
    parts = text.split("\n")
    for i, part in enumerate(parts):
        if i > 0:
            result.append(("\n", {}))
        if part:
            result.append((part, opts))


def _parse_runs_candidate(text: object) -> list[object] | None:
    if _is_object_list(text):
        return text

    if isinstance(text, str) and text.startswith("["):
        try:
            parsed: object = json.loads(text)
        except (json.JSONDecodeError, ValueError):
            return None

        if _is_object_list(parsed):
            return parsed

    return None


def _parse_run_item(item: object) -> tuple[str, RunOverridesMap] | None:
    if isinstance(item, str):
        return item, {}

    if _is_object_list(item):
        run_text = str(item[0]) if len(item) > 0 else ""
        opts_raw: object = item[1] if len(item) > 1 else {}
        return run_text, _coerce_overrides(opts_raw)

    if _is_str_object_dict(item) and "text" in item:
        run_text = str(item.get("text", ""))
        opts = {k: v for k, v in item.items() if k != "text"}
        return run_text, opts

    return None


def _parse_plain_text_runs(text: object) -> list[TextRun]:
    text_str = text if isinstance(text, str) else str(text)
    normalized = text_str.replace("\n", "\\n")

    result: list[TextRun] = []
    for i, line in enumerate(normalized.split("\\n")):
        if i > 0:
            result.append(("\n", {}))
        result.append((line, {}))

    return result


def parse_text_arg(text: object) -> list[TextRun]:
    """
    Parse text argument into list of (text, overrides) tuples.

    Formats:
        "Plain text"                                → [("Plain text", {})]
        "Line 1\\nLine 2"                            → with line breaks
        [["Bold ", {"bold": true}], [" normal"]]    → JSON runs
    """
    runs = _parse_runs_candidate(text)
    if runs is None:
        return _parse_plain_text_runs(text)

    result: list[TextRun] = []
    for item in runs:
        parsed_run = _parse_run_item(item)
        if parsed_run is None:
            continue
        run_text, opts = parsed_run
        _expand_newlines(result, run_text, opts)

    return result


def is_paragraphs_format(text_arg: object) -> bool:
    """Check if text_arg is a multi-paragraph spec."""
    if _is_str_object_dict(text_arg) and "paragraphs" in text_arg:
        return True

    if isinstance(text_arg, str) and text_arg.strip().startswith("{"):
        try:
            parsed: object = json.loads(text_arg)
        except (json.JSONDecodeError, ValueError):
            return False

        return _is_str_object_dict(parsed) and "paragraphs" in parsed

    return False


def _to_int(value: object) -> int | None:
    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float, str)):
        try:
            return int(value)
        except ValueError:
            return None

    return None


def _to_float(value: object) -> float | None:
    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    return None


def _parse_paragraph(p_obj: object) -> ParagraphSpec | None:
    if not _is_str_object_dict(p_obj):
        return None

    para: ParagraphSpec = {}

    if "runs" in p_obj:
        para["runs"] = p_obj["runs"]

    level = _to_int(p_obj.get("level"))
    if level is not None:
        para["level"] = level

    alignment_raw = p_obj.get("alignment")
    if isinstance(alignment_raw, str):
        para["alignment"] = alignment_raw

    spacing_before = _to_float(p_obj.get("spacing_before"))
    if spacing_before is not None:
        para["spacing_before"] = spacing_before

    spacing_after = _to_float(p_obj.get("spacing_after"))
    if spacing_after is not None:
        para["spacing_after"] = spacing_after

    line_spacing = _to_float(p_obj.get("line_spacing"))
    if line_spacing is not None:
        para["line_spacing"] = line_spacing

    bullet_raw = p_obj.get("bullet")
    if isinstance(bullet_raw, (bool, str)):
        para["bullet"] = bullet_raw

    return para


def parse_paragraphs_arg(text_arg: object) -> list[ParagraphSpec]:
    """Parse a multi-paragraph spec."""

    parsed: object = text_arg
    if isinstance(parsed, str):
        parsed = json.loads(parsed)

    if not _is_str_object_dict(parsed):
        raise ValueError("paragraphs spec must be an object")

    paragraphs_raw = parsed.get("paragraphs")
    if not _is_object_list(paragraphs_raw):
        raise ValueError("paragraphs spec missing 'paragraphs' list")

    paragraphs: list[ParagraphSpec] = []
    for p_obj in paragraphs_raw:
        para = _parse_paragraph(p_obj)
        if para is not None:
            paragraphs.append(para)

    return paragraphs


def write_to_shape(shape: Shape, text_arg: object) -> None:
    """Write content to a shape, dispatching based on format."""
    from .editor import add_line_break, add_run, snapshot_defaults, write_paragraphs

    if is_paragraphs_format(text_arg):
        write_paragraphs(shape, parse_paragraphs_arg(text_arg))
        return

    runs = parse_text_arg(text_arg)
    tf = shape.text_frame
    defaults = snapshot_defaults(tf)
    defaults_map: dict[str, object] = dict(defaults)

    tf.clear()
    p = tf.paragraphs[0]
    for text, overrides in runs:
        if text == "\n":
            add_line_break(p)
            continue

        merged: dict[str, object] = dict(defaults_map)
        merged.update(overrides)
        add_run(p, text, **merged)


def _preview_fragment_text(fragment: object) -> str:
    if isinstance(fragment, str):
        return fragment

    if _is_str_object_dict(fragment) and "text" in fragment:
        return str(fragment.get("text", ""))

    if _is_object_list(fragment):
        return str(fragment[0]) if len(fragment) > 0 else ""

    return str(fragment)


def _preview_runs_value(runs: object) -> str:
    if isinstance(runs, str):
        return runs

    if _is_object_list(runs):
        return "".join(_preview_fragment_text(fragment) for fragment in runs)

    return str(runs)


def text_preview(text_arg: object) -> str:
    """Compact preview for before/after display."""

    if is_paragraphs_format(text_arg):
        paras = parse_paragraphs_arg(text_arg)
        parts: list[str] = []

        for para in paras:
            level = para.get("level", 0)
            prefix = "  " * level + ("• " if level > 0 else "")
            runs = para.get("runs", "")
            parts.append(f"{prefix}{_preview_runs_value(runs)}")

        return " ¶ ".join(parts)

    parts: list[str] = []
    for text, opts in parse_text_arg(text_arg):
        if text == "\n":
            parts.append("\\n")
        elif opts:
            flags = ",".join(f"{k}={v}" for k, v in opts.items())
            parts.append(f"[{text}|{flags}]")
        else:
            parts.append(text)

    return "".join(parts)
